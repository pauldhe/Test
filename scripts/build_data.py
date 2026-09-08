#!/usr/bin/env python3
"""Build the dashboard's JSON bundle from the Bloomberg vendor workbook.

The workbook itself lives in the private `pauldhe/Data` repository and is never
committed here. Point --workbook at a local checkout of it:

    python scripts/build_data.py --workbook ../Data/data/raw/data_bb.xlsx

Outputs into data/ (git-ignored):

    data/index.json            metadata, date axis, per-series summary + changes
    data/series/<category>.json  full value arrays, lazy-loaded by the dashboard

See docs/DATA.md for the exact shape of each file.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import re
import sys
from pathlib import Path

import openpyxl

SHEET = "D"
ROW_CATEGORY, ROW_NOTATION, ROW_TICKER, ROW_FIELD = 6, 7, 8, 9
ROW_FIRST_DATA = 11  # row 10's date cell is #NAME? and cannot be dated
FIRST_SERIES_COL = 2

NA_MARKERS = frozenset({"#N/A N/A", "#VALUE!", "#NAME?", "#N/A", "N/A", ""})

CATEGORY_CANONICAL = {
    "irs": "IRS", "fx": "FX", "rate": "Rate", "credit": "Credit",
    "equity": "Equity", "commodity": "Commodity", "inflation": "Inflation",
    "etf": "ETF", "macro": "Macro", "other": "Other",
}

# Display order and Korean labels for the dashboard's category tabs.
CATEGORY_ORDER = ["Rate", "IRS", "Credit", "FX", "Inflation", "Equity", "ETF", "Commodity", "Macro", "Other"]
CATEGORY_LABEL = {
    "Rate": "금리", "IRS": "IRS", "Credit": "크레딧", "FX": "환율",
    "Inflation": "인플레이션", "Equity": "주식", "ETF": "ETF",
    "Commodity": "원자재", "Macro": "매크로", "Other": "기타",
}

_UNIT_BY_KEYWORD = (
    ("CDS", "bp"), ("변동성지수", "pt"), ("NonFarm_Payrolls", "k"),
    ("재고증감", "bn"), ("실업률", "%"), ("CFNAI", "pt"), ("gap", "%"),
)
_UNIT_BY_CATEGORY = {
    "IRS": "%", "Rate": "%", "Credit": "%", "Inflation": "%", "Macro": "%",
    "Equity": "idx", "ETF": "idx", "Commodity": "idx", "Other": "idx", "FX": "",
}

# Government curve points, for the yield-curve view. Base rates, linker yields
# and overnight benchmarks are not curve points and are excluded.
CURVE_COUNTRIES = ["미국", "한국", "유로", "일본", "캐나다", "호주", "영국"]
_CURVE_RE = re.compile(r"^(" + "|".join(CURVE_COUNTRIES) + r")_(\d+(?:\.\d+)?)(m|y)$")

# Lookback windows for the change columns, in calendar days. YTD is handled
# separately because its length depends on the as-of date.
CHANGE_WINDOWS = {"1d": 1, "1w": 7, "1m": 30, "3m": 91, "1y": 365}


def clean(cell: object) -> str:
    return str(cell).strip() if cell is not None else ""


def slug(text: str) -> str:
    """Series id from the notation, preserving rating notches.

    ``+``/``-`` are spelled out before punctuation is stripped, otherwise
    ``AA+`` and ``AA-`` collapse onto one id and two distinct credit curves get
    silently merged.
    """
    s = text.strip().replace("&", "and").replace("+", "_plus_").replace("-", "_minus_")
    s = re.sub(r"[^\w가-힣]+", "_", s, flags=re.UNICODE)
    return re.sub(r"_+", "_", s).strip("_") or "series"


def to_number(cell: object) -> float | None:
    if cell is None or isinstance(cell, bool):
        return None
    if isinstance(cell, (int, float)):
        value = float(cell)
        return value if math.isfinite(value) else None
    text = str(cell).strip()
    if text in NA_MARKERS or text.startswith("#"):
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def resolve_unit(category: str, notation: str) -> str:
    for keyword, unit in _UNIT_BY_KEYWORD:
        if keyword.lower() in notation.lower():
            return unit
    return _UNIT_BY_CATEGORY.get(category, "")


def parse_curve_point(category: str, notation: str) -> tuple[str, float] | None:
    """Return (country, tenor_in_years) when the notation is a curve point."""
    if category != "Rate":
        return None
    match = _CURVE_RE.match(notation)
    if not match:
        return None
    country, size, unit = match.groups()
    years = float(size) / 12 if unit == "m" else float(size)
    return country, years


def assign_ids(rows: list[tuple[int, str, str]]) -> dict[int, str]:
    by_slug: dict[str, list[tuple[int, str, str]]] = {}
    for column, notation, ticker in rows:
        by_slug.setdefault(slug(notation), []).append((column, notation, ticker))

    ids: dict[int, str] = {}
    for base, members in by_slug.items():
        if len(members) == 1:
            ids[members[0][0]] = base
            continue
        seen: set[str] = set()
        for ordinal, (column, _notation, ticker) in enumerate(members, start=1):
            token = slug(ticker.split()[0]) if ticker.split() else ""
            candidate = f"{base}__{token}" if token else f"{base}__{ordinal}"
            if candidate in seen:
                candidate = f"{candidate}_{ordinal}"
            seen.add(candidate)
            ids[column] = candidate
    return ids


def load_workbook(path: Path) -> tuple[list[dt.date], list[dict]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise SystemExit(f"sheet {SHEET!r} not found in {path}")
    sheet = workbook[SHEET]

    header = {i: r for i, r in enumerate(sheet.iter_rows(min_row=1, max_row=ROW_FIELD, values_only=True), 1)}

    def cell(row: int, column: int) -> str:
        values = header.get(row, ())
        return clean(values[column - 1]) if column - 1 < len(values) else ""

    width = len(header.get(ROW_NOTATION, ()))
    columns = [c for c in range(FIRST_SERIES_COL, width + 1) if cell(ROW_NOTATION, c)]
    if not columns:
        raise SystemExit("no series columns found; the workbook layout may have changed")

    ids = assign_ids([(c, cell(ROW_NOTATION, c), cell(ROW_TICKER, c)) for c in columns])

    series: list[dict] = []
    for column in columns:
        raw_category = cell(ROW_CATEGORY, column)
        category = CATEGORY_CANONICAL.get(raw_category.lower(), raw_category or "Other")
        notation = cell(ROW_NOTATION, column)
        curve = parse_curve_point(category, notation)
        series.append({
            "id": ids[column], "col": column, "cat": category, "notation": notation,
            "ticker": cell(ROW_TICKER, column), "field": cell(ROW_FIELD, column),
            "unit": resolve_unit(category, notation),
            "curve": {"country": curve[0], "tenor": curve[1]} if curve else None,
            "values": [],
        })

    dates: list[dt.date] = []
    for row in sheet.iter_rows(min_row=ROW_FIRST_DATA, values_only=True):
        stamp = row[0] if row else None
        if not isinstance(stamp, dt.datetime):
            if stamp is None:
                break
            continue
        dates.append(stamp.date())
        for item in series:
            index = item["col"] - 1
            item["values"].append(to_number(row[index]) if index < len(row) else None)

    workbook.close()
    return dates, series


def last_value_at_or_before(values: list[float | None], index: int) -> float | None:
    for i in range(min(index, len(values) - 1), -1, -1):
        if values[i] is not None:
            return values[i]
    return None


def index_at_or_before(dates: list[dt.date], target: dt.date) -> int | None:
    lo, hi, found = 0, len(dates) - 1, None
    while lo <= hi:
        mid = (lo + hi) // 2
        if dates[mid] <= target:
            found, lo = mid, mid + 1
        else:
            hi = mid - 1
    return found


def change_convention(category: str, unit: str) -> tuple[str, str]:
    """How this series' change should be measured and labelled.

    Three conventions, because one does not fit everything:

    * ``bp``  - yields and spreads move in basis points.
    * ``abs`` - an absolute difference, for series already denominated in
      something meaningful: macro rates in percentage points (GDP growth going
      from -13.8% to 7.9% is +21.7%p, not +2170bp), index points, thousands of
      payrolls.
    * ``pct`` - percent change, for levels and prices.

    The percent change of a percentage yield is never computed; it is not a
    meaningful quantity.
    """
    if unit == "%":
        return ("abs", "%p") if category == "Macro" else ("bp", "bp")
    if unit == "bp":
        return ("bp", "bp")
    if unit in ("pt", "k", "bn"):
        return ("abs", unit)
    return ("pct", "%")


def summarise(item: dict, dates: list[dt.date]) -> dict:
    """Latest level plus change over each lookback window."""
    values = item["values"]
    populated = [i for i, v in enumerate(values) if v is not None]
    if not populated:
        return {"n": 0, "last": None, "lastDate": None, "changes": {}, "mode": None,
                "min": None, "max": None, "first": None, "firstDate": None}

    first_index, last_index = populated[0], populated[-1]
    latest = values[last_index]
    as_of = dates[last_index]
    mode, suffix = change_convention(item["cat"], item["unit"])

    changes: dict[str, float | None] = {}
    targets = {name: as_of - dt.timedelta(days=days) for name, days in CHANGE_WINDOWS.items()}
    targets["ytd"] = dt.date(as_of.year, 1, 1) - dt.timedelta(days=1)

    for name, target in targets.items():
        anchor = index_at_or_before(dates, target)
        previous = last_value_at_or_before(values, anchor) if anchor is not None else None
        if previous is None or anchor is None or anchor < first_index:
            changes[name] = None
        elif mode == "bp":
            changes[name] = round((latest - previous) * (100 if item["unit"] == "%" else 1), 4)
        elif mode == "abs":
            changes[name] = round(latest - previous, 4)
        else:
            changes[name] = round((latest / previous - 1) * 100, 4) if previous else None

    numeric = [v for v in values if v is not None]
    return {
        "n": len(numeric), "last": latest, "lastDate": as_of.isoformat(),
        "first": values[first_index], "firstDate": dates[first_index].isoformat(),
        "min": min(numeric), "max": max(numeric), "changes": changes,
        "mode": mode, "suffix": suffix,
        "spark": sparkline(values, first_index, last_index),
    }


def sparkline(values: list[float | None], first_index: int, last_index: int,
              points: int = 48) -> list[float | None]:
    """Evenly-spaced sample of roughly the last year, for the overview table.

    Carried in index.json so the overview renders without pulling the multi-MB
    value files.
    """
    start = max(first_index, last_index - 260)
    if last_index <= start:
        return []
    step = (last_index - start) / (points - 1)
    sampled: list[float | None] = []
    for k in range(points):
        i = min(last_index, start + int(round(k * step)))
        value = next((values[j] for j in range(i, start - 1, -1) if values[j] is not None), None)
        sampled.append(None if value is None else round(value, 6))
    return sampled


def mark_duplicates(series: list[dict]) -> None:
    """Tag series whose values are identical to an earlier column's.

    The vendor sheet repeats several series -- the same ticker pasted twice, or
    one ticker filed under two notations. They are kept (the workbook is stored
    as delivered) but flagged, so a chart does not silently draw the same line
    twice under two names.
    """
    first_seen: dict[tuple, str] = {}
    for item in series:
        item["dupeOf"] = None
        if item["summary"]["n"] == 0:
            continue
        key = (item["field"], tuple(item["values"]))
        if key in first_seen:
            item["dupeOf"] = first_seen[key]
        else:
            first_seen[key] = item["id"]


def stale_tail_diagnostic(series: list[dict], dates: list[dt.date], lookback: int = 40) -> dict:
    """Flag a final row that merely repeats the previous day's values.

    An extract taken before the close leaves most tickers carrying their prior
    close, which would make every "1-day change" on the dashboard read zero and
    date the snapshot a day late. A genuine trading day still repeats some
    values -- step-held macro series and unchanged policy rates -- so the test
    is against the recent baseline rate, not against zero.
    """
    def identical_rate(i: int, j: int) -> float | None:
        both = same = 0
        for item in series:
            a, b = item["values"][i], item["values"][j]
            if a is not None and b is not None:
                both += 1
                same += a == b
        return same / both if both else None

    if len(dates) < lookback + 2:
        return {"suspect": False}

    last = identical_rate(len(dates) - 1, len(dates) - 2)
    baseline = [
        r for r in (identical_rate(i, i - 1) for i in range(len(dates) - lookback - 1, len(dates) - 1))
        if r is not None
    ]
    if last is None or not baseline:
        return {"suspect": False}

    baseline.sort()
    median = baseline[len(baseline) // 2]
    suspect = last > 0.5 and last > 2 * median

    return {
        "suspect": suspect,
        "lastDate": dates[-1].isoformat(),
        "previousDate": dates[-2].isoformat(),
        "identicalRate": round(last, 4),
        "baselineRate": round(median, 4),
    }


def main(argv: list[str] | None = None) -> int:
    repo = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", type=Path, default=repo.parent / "Data/data/raw/data_bb.xlsx",
                        help="path to data_bb.xlsx from the private Data repository")
    parser.add_argument("--out", type=Path, default=repo / "data")
    parser.add_argument("--precision", type=int, default=6)
    args = parser.parse_args(argv)

    if not args.workbook.exists():
        parser.error(
            f"workbook not found: {args.workbook}\n"
            "Clone the private Data repository next to this one, or pass --workbook."
        )

    dates, series = load_workbook(args.workbook)
    print(f"parsed {len(series)} series x {len(dates)} dates "
          f"({dates[0]} .. {dates[-1]})", file=sys.stderr)

    for item in series:
        item["summary"] = summarise(item, dates)

    mark_duplicates(series)

    categories = [c for c in CATEGORY_ORDER if any(s["cat"] == c for s in series)]
    categories += sorted({s["cat"] for s in series} - set(categories))

    series_dir = args.out / "series"
    series_dir.mkdir(parents=True, exist_ok=True)

    index = {
        "generated": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
        "source": args.workbook.name,
        "demo": "demo" in args.workbook.name.lower(),
        "dates": [d.isoformat() for d in dates],
        "categories": [
            {"key": c, "label": CATEGORY_LABEL.get(c, c),
             "count": sum(1 for s in series if s["cat"] == c)}
            for c in categories
        ],
        "series": [
            {k: item[k] for k in ("id", "cat", "notation", "ticker", "field", "unit", "curve", "dupeOf")}
            | {"summary": item["summary"]}
            for item in series
        ],
        "totals": {
            "series": len(series), "dates": len(dates),
            "observations": sum(s["summary"]["n"] for s in series),
            "empty": [s["id"] for s in series if s["summary"]["n"] == 0],
        },
        "diagnostics": {
            "staleTail": stale_tail_diagnostic(series, dates),
            "duplicates": [
                {"id": s["id"], "of": s["dupeOf"]} for s in series if s["dupeOf"]
            ],
        },
    }
    (args.out / "index.json").write_text(
        json.dumps(index, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    for category in categories:
        payload = {
            item["id"]: [None if v is None else round(v, args.precision) for v in item["values"]]
            for item in series if item["cat"] == category
        }
        (series_dir / f"{category}.json").write_text(
            json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    written = sum(p.stat().st_size for p in [args.out / "index.json", *series_dir.glob("*.json")])
    print(f"wrote {args.out}/index.json + {len(categories)} category files "
          f"({written / 1e6:.1f} MB), {index['totals']['observations']:,} observations",
          file=sys.stderr)
    stale = index["diagnostics"]["staleTail"]
    if stale.get("suspect"):
        print(f"warning: last row {stale['lastDate']} repeats {stale['identicalRate']:.0%} of "
              f"{stale['previousDate']}'s values (recent norm {stale['baselineRate']:.0%}); "
              "it looks like a pre-close snapshot", file=sys.stderr)
    if index["totals"]["empty"]:
        print(f"note: {len(index['totals']['empty'])} series have no data: "
              f"{', '.join(index['totals']['empty'])}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
